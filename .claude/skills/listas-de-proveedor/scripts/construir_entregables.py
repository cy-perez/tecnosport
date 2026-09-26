#!/usr/bin/env python3
"""
Arma los entregables finales a partir de productos.json ya enriquecido.

Uso:
    python3 construir_entregables.py productos.json --salida entregables/

Espera que cada producto de productos.json ya tenga:
    titulo, descripcion, precio_proveedor_cop, precio_mercado_cop,
    colores_oficiales, fuentes_precio

Produce:
    entregables/catalogo-<fecha>.zip       una carpeta por producto con su .txt
    entregables/comparativo-<fecha>.xlsx   título / precio lista / precio mercado / ganancia

Las fotos quedaron fuera del flujo de la skill (regla 16). `--imagenes` sigue
existiendo para el caso en que ya haya un lote retocado a mano con
`fotos-estudio-degradado`: sin la bandera el ZIP sale solo con las fichas y no
se emite ningún FOTOS-PENDIENTES.md.
"""

import argparse
import json
import re
import shutil
import zipfile
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

EXT_IMAGEN = (".jpg", ".jpeg", ".png", ".webp", ".avif")


def ficha_txt(p: dict) -> str:
    colores = ", ".join(p.get("colores_oficiales") or p.get("colores_familia") or []) or "único"
    precio = p.get("precio_mercado_cop")
    L = [
        f"TÍTULO: {p['titulo']}",
        f"CATEGORÍA: {p.get('categoria', '')}",
        f"MARCA: {p.get('marca') or 'por confirmar'}",
        f"CONDICIÓN: {p.get('condicion', '')}",
        f"MEMORIA: {p.get('ram') or 'n/a'} RAM"
        + (f" (+{p['ram_virtual']} virtual)" if p.get("ram_virtual") else "")
        + f" / {p.get('almacenamiento') or 'n/a'}",
        f"RED: {p.get('red') or 'n/a'}",
        f"COLORES: {colores}",
        f"PRECIO DE PUBLICACIÓN (COP): {precio:,}".replace(",", ".") if precio else "PRECIO DE PUBLICACIÓN (COP): por definir",
        f"SKU SUGERIDO: {p.get('id', '')}",
        "",
        f"META TÍTULO: {p.get('meta_titulo', '')}",
        f"META DESCRIPCIÓN: {p.get('meta_descripcion', '')}",
        "",
        "DESCRIPCIÓN",
        "-" * 60,
        p.get("descripcion") or "(pendiente de redactar)",
    ]
    if p.get("supuestos"):
        L += ["", "SUPUESTOS APLICADOS AL LEER LA LISTA", "-" * 60]
        L += [f"- {a}" for a in p["supuestos"]]
    if p.get("revisar"):
        L += ["", "PENDIENTES ANTES DE PUBLICAR", "-" * 60]
        L += [f"- {r}" for r in p["revisar"]]
    if p.get("fuentes_precio"):
        L += ["", "FUENTES DEL PRECIO DE MERCADO", "-" * 60]
        for f in p["fuentes_precio"]:
            # el parser y asignar_precios.py escriben precio_cop y enlace; se
            # aceptan también precio y url, que es como lo escribe una persona
            # cuando completa una fuente a mano
            precio = f.get("precio_cop") or f.get("precio")
            enlace = f.get("enlace") or f.get("url") or ""
            nivel = f" ({f['nivel']})" if f.get("nivel") else ""
            monto = f"{precio:,}".replace(",", ".") if precio else "?"
            L.append(f"- {f.get('tienda', '?')}{nivel}: {monto} — {enlace}")
    return "\n".join(L) + "\n"


def construir(datos: dict, dir_imagenes: Path, salida: Path):
    fecha = datos.get("fecha_lista") or date.today().isoformat()
    salida.mkdir(parents=True, exist_ok=True)
    staging = salida / f"catalogo-{fecha}"
    if staging.exists():
        shutil.rmtree(staging)
    staging.mkdir(parents=True)

    sin_fotos = []
    for p in datos["productos"]:
        carpeta = staging / p["id"]
        carpeta.mkdir(parents=True, exist_ok=True)
        (carpeta / f"{p['id']}.txt").write_text(ficha_txt(p), encoding="utf-8")

        # Las fotos salieron del flujo (regla 16). Sin `--imagenes` no hay nada
        # que copiar y tampoco tiene sentido avisar que faltan: faltarían en
        # todos los productos, siempre. Con la bandera puesta —apuntando a un
        # lote ya retocado a mano— el comportamiento es el de antes.
        if not dir_imagenes:
            continue

        origen = dir_imagenes / p["id"]
        todas = []
        if origen.is_dir():
            # `fotos-estudio-degradado` agrupa por producto y separa por ancho
            # (`<producto>/maestra/`, `1200/`, `800/`…) en cuanto las fotos le
            # llegan en subcarpetas, que es como sale de ella. Para el ZIP
            # se lleva la maestra, que es la mejor versión de cada toma.
            maestra = origen / "maestra"
            if maestra.is_dir():
                todas = sorted(f for f in maestra.iterdir()
                               if f.suffix.lower() in EXT_IMAGEN)
            else:
                todas = sorted(f for f in origen.iterdir()
                               if f.suffix.lower() in EXT_IMAGEN)

        # Las maestras se llaman <id>-01.jpg; el resto son variantes responsive.
        patron = re.compile(rf"^{re.escape(p['id'])}-\d{{2}}$")
        maestras = [f for f in todas if patron.match(f.stem)]
        if maestras:
            for f in maestras[:4] + [f for f in todas if f not in maestras]:
                shutil.copy2(f, carpeta / f.name)
            n_fotos = len(maestras)
        else:
            for i, f in enumerate(todas[:4], start=1):
                shutil.copy2(f, carpeta / f"{p['id']}-{i:02d}{f.suffix.lower()}")
            n_fotos = len(todas)

        if n_fotos < 4:
            sin_fotos.append((p["titulo"], n_fotos))
            pendiente = [
                f"# Fotos pendientes — {p['titulo']}",
                "",
                f"Hay {n_fotos} de 4 fotos. Faltan {4 - n_fotos}.",
                "",
                "Orden de las fotos: 01 frontal · 02 posterior · 03 ángulo o lateral · 04 detalle.",
                "Fuentes admitidas: el paquete de imágenes del proveedor, el portal de",
                "partners si la tienda es revendedor autorizado, Open Icecat o fotos propias.",
                "Registrar el origen de cada una antes de publicar.",
                "",
                "NO sirven las salas de prensa del fabricante (Apple Newsroom, Samsung",
                "Mobile Press y equivalentes): sus condiciones autorizan uso editorial o",
                "personal, no publicar el producto en una tienda.",
                "",
            ]
            for c in p.get("imagenes_candidatas", []):
                pendiente.append(f"- {c}")
            (carpeta / "FOTOS-PENDIENTES.md").write_text("\n".join(pendiente), encoding="utf-8")

    ruta_zip = salida / f"catalogo-{fecha}.zip"
    with zipfile.ZipFile(ruta_zip, "w", zipfile.ZIP_DEFLATED) as z:
        for f in sorted(staging.rglob("*")):
            if f.is_file():
                z.write(f, f.relative_to(staging.parent))

    ruta_xlsx = salida / f"comparativo-{fecha}.xlsx"
    escribir_excel(datos, ruta_xlsx, fecha)
    shutil.rmtree(staging)
    return ruta_zip, ruta_xlsx, sin_fotos


def escribir_excel(datos: dict, ruta: Path, fecha: str):
    wb = Workbook()
    ws = wb.active
    ws.title = "Comparativo"
    encabezados = [
        "Título del producto",
        "Precio lista proveedor (COP)",
        "Precio promedio mercado Colombia (COP)",
        "Ganancia (COP)",
    ]
    ws.append(encabezados)
    relleno = PatternFill("solid", fgColor="1F3A5F")
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = relleno
        c.alignment = Alignment(vertical="center", wrap_text=True)

    for p in datos["productos"]:
        costo = p.get("precio_proveedor_cop")
        mercado = p.get("precio_mercado_cop")
        ganancia = (mercado - costo) if (costo and mercado) else None
        ws.append([p["titulo"], costo, mercado, ganancia])

    for fila in ws.iter_rows(min_row=2, min_col=2, max_col=4):
        for c in fila:
            c.number_format = '#,##0'
    for col, ancho in zip("ABCD", (52, 26, 32, 20)):
        ws.column_dimensions[col].width = ancho
    ws.freeze_panes = "A2"

    ws2 = wb.create_sheet("Detalle")
    ws2.append(["Título", "Categoría", "Marca", "Condición", "Colores",
                "Margen %", "Fuentes del precio", "Pendientes", "Supuestos"])
    for c in ws2[1]:
        c.font = Font(bold=True)
    for p in datos["productos"]:
        costo, mercado = p.get("precio_proveedor_cop"), p.get("precio_mercado_cop")
        margen = round((mercado - costo) / mercado * 100, 1) if (costo and mercado) else None
        ws2.append([
            p["titulo"], p.get("categoria"), p.get("marca"), p.get("condicion"),
            ", ".join(p.get("colores_oficiales") or p.get("colores_familia") or []),
            margen,
            " | ".join(f.get("tienda", "") for f in p.get("fuentes_precio", [])),
            " | ".join(p.get("revisar", [])),
            " | ".join(p.get("supuestos", [])),
        ])
    for i, ancho in enumerate((46, 16, 14, 14, 26, 10, 40, 60, 50), start=1):
        ws2.column_dimensions[get_column_letter(i)].width = ancho

    ws3 = wb.create_sheet("Descartados")
    ws3.append(["Texto original", "Motivo"])
    for c in ws3[1]:
        c.font = Font(bold=True)
    for p in datos.get("descartados", []):
        ws3.append([p.get("texto_origen", ""), p.get("motivo", "")])
    ws3.column_dimensions["A"].width = 50
    ws3.column_dimensions["B"].width = 40

    wb.save(ruta)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("productos")
    ap.add_argument("--imagenes", default=None,
                    help="carpeta de un lote ya retocado con fotos-estudio-degradado; "
                         "sin ella el ZIP sale solo con las fichas")
    ap.add_argument("--salida", default="entregables")
    args = ap.parse_args()

    datos = json.loads(Path(args.productos).read_text(encoding="utf-8"))
    zipf, xlsx, sin_fotos = construir(
        datos, Path(args.imagenes) if args.imagenes else None, Path(args.salida)
    )
    print(f"ZIP  -> {zipf}")
    print(f"XLSX -> {xlsx}")
    if sin_fotos:
        print(f"\n{len(sin_fotos)} productos con menos de 4 fotos:")
        for titulo, n in sin_fotos[:15]:
            print(f"  - {titulo}: {n}/4")


if __name__ == "__main__":
    main()
